#!/usr/bin/env python3
"""Create the Chinese Gkeyll/FNO reproduction and alignment report."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/reports/GKEYLL_FNO_REPRODUCTION_ALIGNMENT_REPORT_zh-CN.docx"

ROUND8 = ROOT / "results/gkeyll_round8/final_report"
ROUND9 = ROOT / "results/gkeyll_round9/final_report"

IMAGES = {
    "field": ROUND8 / "paper_field_energy_truth_hp_fno.png",
    "moments": ROUND8 / "paper_moments_truth_hp_fno.png",
    "moment_errors": ROUND8 / "paper_moment_errors_fno_hp.png",
    "closure": ROUND8 / "paper_closure_truth_hp_fno.png",
    "stability": ROUND8 / "stability_constraint_seed_comparison.png",
    "temporal": ROUND9 / "single_case_temporal_extrapolation.png",
    "cross_case": ROUND9 / "cross_case_generalization.png",
    "protocol_gap": ROUND9 / "protocol_generalization_gap.png",
}

NAVY = "17365D"
BLUE = "2F75B5"
TEAL = "0F6B78"
LIGHT_BLUE = "DDEBF7"
LIGHT_TEAL = "DDEFEF"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9E2F3"
TEXT_GRAY = RGBColor(90, 98, 108)
RED = "C00000"
AMBER = "FFF2CC"
GREEN = "E2F0D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=100, bottom=85, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_header_table_style(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for i, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(8.6)
                    set_run_font(run)
        if i == 0:
            for cell in row.cells:
                set_cell_shading(cell, NAVY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        elif i % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_GRAY)


def set_run_font(run, east_asia="Microsoft YaHei", latin="Aptos") -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def set_style_font(style, size, bold=False, color=None, east_asia="Microsoft YaHei") -> None:
    style.font.name = "Aptos"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, 10.2)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    set_style_font(styles["Title"], 27, True, NAVY)
    styles["Title"].paragraph_format.space_after = Pt(16)
    set_style_font(styles["Subtitle"], 13, False, TEAL)

    for name, size, color in (
        ("Heading 1", 18, NAVY),
        ("Heading 2", 14, TEAL),
        ("Heading 3", 11.5, BLUE),
    ):
        style = styles[name]
        set_style_font(style, size, True, color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)

    for style_name in ("Caption", "Quote"):
        set_style_font(styles[style_name], 9, False, None)
    styles["Caption"].font.italic = False
    styles["Caption"].font.color.rgb = TEXT_GRAY
    styles["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Quote"].font.color.rgb = RGBColor.from_string(NAVY)
    styles["Quote"].paragraph_format.left_indent = Cm(0.6)
    styles["Quote"].paragraph_format.right_indent = Cm(0.4)


def set_default_language(doc: Document) -> None:
    styles = doc.styles
    for style in styles:
        if not hasattr(style, "_element") or style._element.rPr is None:
            continue
        lang = style._element.rPr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            style._element.rPr.append(lang)
        lang.set(qn("w:eastAsia"), "zh-CN")
        lang.set(qn("w:val"), "zh-CN")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        section.header_distance = Cm(0.7)
        section.footer_distance = Cm(0.7)
        header = section.header.paragraphs[0]
        header.text = "Huang et al. (2025) FNO 热流闭合复现与物理合同对齐"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            set_run_font(run)
            run.font.size = Pt(8)
            run.font.color.rgb = TEXT_GRAY
        add_page_number(section.footer.paragraphs[0])


def add_horizontal_rule(paragraph, color=NAVY, size="12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(38)
    add_horizontal_rule(p, TEAL, "28")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Huang et al. (2025)\n非线性 Landau 阻尼 FNO 热流闭合复现报告")
    set_run_font(run)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("物理合同对齐 · 同轨迹闭环复现 · 严格时间外推 · 跨案例泛化")

    doc.add_paragraph("\n")
    box = doc.add_table(rows=3, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.autofit = False
    box.columns[0].width = Cm(14.5)
    texts = [
        "复现对象：Machine-learning heat flux closure for multi-moment fluid modeling of nonlinear Landau damping",
        "复现主案例：1D1V Gkeyll，k = 0.35，A = 0.10，t = 0–40",
        "报告定位：内部技术报告；区分论文同轨迹复现能力与样本外泛化能力",
    ]
    for i, text in enumerate(texts):
        cell = box.cell(i, 0)
        set_cell_shading(cell, LIGHT_BLUE if i != 2 else LIGHT_TEAL)
        set_cell_margins(cell, 170, 200, 170, 200)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r)
        r.font.size = Pt(10.5)
        if i == 2:
            r.font.bold = True
            r.font.color.rgb = RGBColor.from_string(TEAL)

    doc.add_paragraph("\n\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("landau-damping-surrogate-standardized 项目组")
    set_run_font(r)
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2026 年 8 月 29 日")
    set_run_font(r)
    r.font.size = Pt(10.5)
    r.font.color.rgb = TEXT_GRAY
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    if level == 1:
        add_horizontal_rule(p, BLUE, "8")


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    r = p.add_run(text)
    set_run_font(r)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_run_font(r)


def add_callout(doc: Document, title: str, body: str, fill=LIGHT_BLUE, title_color=NAVY) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 150, 180, 150, 180)
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    set_run_font(r)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(title_color)
    r.font.size = Pt(11)
    r = p.add_run(body)
    set_run_font(r)
    r.font.size = Pt(10)
    p.paragraph_format.line_spacing = 1.2
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc: Document, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = str(header)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    set_repeat_header_table_style(table)
    return table


def add_picture(doc: Document, path: Path, caption: str, width_inches=6.7, source_note=None) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width_inches))
    cap = doc.add_paragraph(style="Caption")
    cap.add_run(caption)
    if source_note:
        src = doc.add_paragraph(style="Caption")
        src.paragraph_format.space_after = Pt(7)
        run = src.add_run(f"数据与绘图源：{source_note}")
        run.font.size = Pt(7.5)


def add_contents(doc: Document) -> None:
    add_heading(doc, "报告导读", 1)
    rows = [
        ("1", "执行摘要", "复现达成度、核心数字、结论边界"),
        ("2", "复现对象与物理合同", "逐项比较论文与我们的物理、数值、数据和 ML 设置"),
        ("3", "从 PIC 初试到论文颗粒度对齐", "解释早期问题、诊断证据和修正路径"),
        ("4", "论文同轨迹复现结果", "Gkeyll 真值、HP 闭合与纯监督 FNO 的闭环比较"),
        ("5", "稳定性和随机种子", "为什么最佳曲线不能代表全部训练运行"),
        ("6", "严格泛化实验", "单案例时间外推与整案例留出"),
        ("7", "结论、价值与下一步", "可宣称成果、不可宣称内容和后续路线"),
    ]
    add_table(doc, ["章节", "主题", "回答的问题"], rows, [1.2, 4.4, 10.7])
    doc.add_paragraph()
    add_callout(
        doc,
        "阅读口径",
        "“论文同轨迹复现”表示模型训练时看过该案例 t=0–40 的真值状态，再在相同初值下自由闭环推进；"
        "它检验闭合映射能否被流体求解器正确调用，但不等价于对未见时间或未见物理参数的泛化。",
        AMBER,
        NAVY,
    )
    doc.add_page_break()


def add_executive_summary(doc: Document) -> None:
    add_heading(doc, "1. 执行摘要", 1)
    add_callout(
        doc,
        "总判断：核心物理现象已在论文同案例、同轨迹协议下复现；通用闭合尚未实现。",
        "对 k=0.35、A=0.10 的 Gkeyll 非线性 Landau 阻尼案例，纯监督 FNO 闭合能够在三矩流体方程中复现"
        "场能先阻尼、达到最低点、随后非线性回升及振荡的主要幅值与相位；标准 HP 闭合则持续过阻尼。"
        "但在未见后期状态或未见整条参数轨迹上，长期闭环误差明显增大。",
        GREEN,
        TEAL,
    )

    add_heading(doc, "1.1 复现达到什么程度", 2)
    rows = [
        ("同轨迹最佳纯监督 FNO", "场能 log10 RMSE", "0.03791", "Gkeyll 真值几乎重合"),
        ("同轨迹 HP", "场能 log10 RMSE", "1.1773", "非线性回升和 bounce 频率失真"),
        ("FNO 闭合量", "相对 L2 / 相关系数", "0.00458 / 0.99999", "在真值状态上高度准确"),
        ("FNO 三矩闭环", "n / u / p 相对 L2", "1.86e-4 / 6.25e-3 / 7.55e-4", "无钳位，推进到 t=40"),
        ("同轨迹 3 seed", "场能误差中位数", "0.1788", "仍有显著随机种子差异"),
        ("单案例时间外推", "场能误差中位数", "0.4443", "仅 2/3 seed 完成 t=40"),
        ("整案例留出", "场能误差中位数", "0.6981 / 0.9625", "数值稳定但幅值、相位漂移"),
    ]
    add_table(doc, ["实验口径", "指标", "结果", "含义"], rows, [4.1, 3.7, 3.3, 5.1])

    add_heading(doc, "1.2 这次复现真正证明了什么", 2)
    add_bullet(doc, "物理可行性：从低阶中心矩 {n,u,p} 学习热通量梯度 ∂xq，确实能把动理学非线性反馈注入三矩流体模型。")
    add_bullet(doc, "求解器可行性：Ampère 场推进、primitive 三矩方程、RK4 stage 级调用和真实初始闭合共同组成了可稳定运行的闭环接口。")
    add_bullet(doc, "相对基线价值：在完全相同的流体推进器中，FNO 显著优于标准 HP，说明提升来自数据驱动闭合，而不是换了求解器。")
    add_bullet(doc, "研究边界：目前最好图是同轨迹重建；严格划分结果说明该模型还不能被称为跨状态、跨参数的通用闭合。")

    add_heading(doc, "1.3 相比早期 PIC 路线为什么有质变", 2)
    p = doc.add_paragraph()
    p.add_run(
        "早期 PIC 实验同时混合了粒子噪声、弱非线性覆盖不足、矩定义与目标变量不一致、PDE/场推进差异和绘图窗口问题。"
        "最终复现把这些因素逐项冻结到论文级颗粒度：先用无粒子噪声的 Gkeyll 单案例验证论文命题，再把中心矩、∂xq 标签、"
        "Vlasov–Ampère、RK4 步长、第一步真实 q、FNO 激活函数和完整 8000 帧训练逐项对齐。因此，改善不是单靠“网络更大”，"
        "而是数据合同、物理合同、数值合同和评估合同同时对齐的结果。"
    )
    doc.add_page_break()


def add_contract(doc: Document) -> None:
    add_heading(doc, "2. 复现对象与物理合同对齐", 1)
    p = doc.add_paragraph()
    p.add_run(
        "这里的“物理合同”是指：动理学问题、归一化、初边值条件、矩定义、闭合目标、流体方程、场推进、时间积分、数据采样和模型调用方式的完整约定。"
        "这些约定中任何一项错位，都可能让离线预测看似合理、闭环推进却产生完全不同的物理。"
    )

    add_heading(doc, "2.1 物理、数据与数值合同逐项对照", 2)
    rows = [
        ("物理维度", "1D1V 无碰撞电子 Landau 阻尼", "相同", "完全对齐"),
        ("离子处理", "固定中和离子背景", "相同", "完全对齐"),
        ("动理学求解器", "Gkeyll Vlasov–Ampère", "Gkeyll Vlasov–Ampère", "完全对齐"),
        ("初始密度", "nₑ=n₀[1+A cos(kx)]", "相同", "完全对齐"),
        ("非线性锚点", "k=0.35，A=0.10", "相同", "完全对齐"),
        ("空间域/边界", "0<x<2π/k，周期边界", "相同", "完全对齐"),
        ("相空间网格", "Nx=64；Nv=64", "Nx=64；Nv=64", "完全对齐"),
        ("速度域", "−6vth 至 +6vth", "相同", "完全对齐"),
        ("动理学时长", "t=0–40", "相同", "完全对齐"),
        ("训练帧", "Δt=0.005，共 8000 帧", "相同", "完全对齐"),
        ("低阶矩", "中心矩 n、u、p", "相同", "完全对齐"),
        ("高阶矩", "q=m∫(v−u)³f dv", "相同", "完全对齐"),
        ("学习目标", "由 {n,u,p} 预测 ∂q/∂x", "相同", "完全对齐"),
        ("流体系统", "三矩 primitive 方程", "相同", "完全对齐"),
        ("电场推进", "Vlasov–Ampère 对应的 Ampère 更新", "相同", "完全对齐"),
        ("闭环积分", "RK4，Δt=0.002，共 20000 步", "相同", "完全对齐"),
        ("第一步闭合", "使用给定初始 q；后续用 FNO", "使用真值初始闭合；后续用 FNO", "完全对齐"),
        ("FNO 主结构", "4 个 Fourier layer，ReLU，频域+实域双通路", "4 层 FNO，ReLU，无 GroupNorm", "结构对齐"),
        ("FNO width/modes", "正文未披露", "训练 width=128、modes=32", "工程选择；无法逐项核对"),
        ("部署谱截断", "正文未披露", "最大 mode=8", "额外稳定化选择"),
        ("主结果训练协议", "同一 k=0.35,A=0.10 全轨迹训练后闭环", "相同", "协议对齐，但不是样本外"),
        ("严格泛化", "正文未提供整案例留出统计", "补做时间外推与 case-wise 留出", "我们的扩展评估"),
    ]
    table = add_table(doc, ["合同项", "论文设置", "我们的最终设置", "对齐结论"], rows, [3.0, 5.1, 5.2, 3.2])
    for i, row in enumerate(table.rows[1:], start=1):
        status = row.cells[3].text
        if "完全" in status:
            set_cell_shading(row.cells[3], GREEN)
        elif "无法" in status or "额外" in status:
            set_cell_shading(row.cells[3], AMBER)
        else:
            set_cell_shading(row.cells[3], LIGHT_BLUE)

    add_heading(doc, "2.2 对齐后的闭合方程", 2)
    equations = [
        "动理学方程：∂tf + v∂xf − E∂vf = 0（电子归一化符号约定）",
        "中心矩：n=∫f dv；u=n⁻¹∫vf dv；p=m∫(v−u)²f dv；q=m∫(v−u)³f dv",
        "连续性：∂tn + ∂x(nu) = 0",
        "速度：∂tu + u∂xu + n⁻¹∂xp = −E",
        "压力：∂tp + u∂xp + 3p∂xu + ∂xq = 0",
        "Ampère：∂tE = nu − ⟨nu⟩",
        "机器学习闭合：FNO[n(x),u(x),p(x)] → ∂xq(x)",
    ]
    for eq in equations:
        p = doc.add_paragraph(style="Quote")
        r = p.add_run(eq)
        set_run_font(r, east_asia="Cambria Math", latin="Cambria Math")
        r.font.size = Pt(10.2)

    add_callout(
        doc,
        "关键合同不是“预测 q 还是预测 ∂xq”的文字差别",
        "压力方程真正使用的是 ∂xq。若用含 PIC 高频噪声的 q 再差分，噪声会被导数放大；若把原始三阶矩 M3 当成中心热通量 q，"
        "还会引入 3uM2 和 2nu³ 项的系统偏差。最终 Gkeyll 复现直接按中心矩构造 q，并以 ∂xq 作为监督和部署目标。",
        LIGHT_TEAL,
        TEAL,
    )
    doc.add_page_break()


def add_alignment_journey(doc: Document) -> None:
    add_heading(doc, "3. 之前遇到的问题，以及如何按论文颗粒度解决", 1)
    p = doc.add_paragraph()
    p.add_run(
        "早期效果差并非一个单独错误，而是多个合同错位叠加。以下表格将“现象—原因—修正—证据”一一对应；其中，"
        "同轨迹复现问题已经基本解决，而样本外闭环反馈仍是未解决的核心问题。"
    )

    rows = [
        ("PIC 与论文 Vlasov 数据不同", "高阶矩噪声大；弱信号下 q 的 seed 噪声可与信号同量级", "先用 Gkeyll 无粒子噪声单案例复现论文；PIC 作为后续域迁移任务", "区分方法复现与 PIC 扩展"),
        ("非线性覆盖不足", "训练帧以线性/弱非线性为主，测试却要求粒子俘获和场能回升", "固定论文非线性锚点 k=.35,A=.10，使用完整 t=0–40、8000 帧", "同轨迹可恢复回升"),
        ("相空间只有条纹", "扰动太弱或时间不足；全速度范围压缩了共振区细节", "使用强非线性案例，并绘制正相速度共振窗口", "Gkeyll 真值可见弯曲/俘获结构"),
        ("矩定义错位", "原始 M2/M3 与中心 p/q 混用", "严格使用中心矩变换 p=M2−nu²；q=M3−3uM2+2nu³", "oracle 残差降至 10⁻⁴ 量级"),
        ("闭合目标错位", "预测 q 后求导可能放大噪声；或直接把 q 喂给压力方程", "对论文主线直接监督 ∂xq，并在 RK stage 调用", "离线闭合 rel-L2 0.00458"),
        ("场方程/PDE 不一致", "Poisson/Ampère、守恒/primitive 形式和电子符号混杂", "对齐 primitive 三矩 + Ampère；用真值 ∂xq 做 oracle 审计", "真值闭合推进场能误差约 10⁻³"),
        ("闭合调用粒度过粗", "每整步调用一次不能代表 RK4 中间状态", "每个 RK stage 重新计算 FNO 闭合；第一步使用真值初始闭合", "20,000 步稳定推进"),
        ("模型结构与论文不同", "GroupNorm/GELU 改变谱特征和闭环敏感性", "去 GroupNorm，使用论文明确写出的 ReLU 与 4 层结构", "Round6 场能误差较旧结构下降 58.1%"),
        ("高模导致闭环不稳定", "高频离线误差虽小，却对 PDE 反馈高度敏感", "扫描部署截断，mode 8 优于 12/16/24", "最佳同轨迹误差继续下降"),
        ("把最好 seed 当整体结论", "训练随机性被漂亮曲线掩盖", "保留最佳图用于论文现象对齐，同时单独报告 3-seed 中位数和范围", "最佳 0.0379；中位数 0.1788"),
        ("训练/测试轨迹重合", "只能证明轨迹重建，不能证明泛化", "增加 t<24→未来外推及完整案例留出", "严格误差上升至 0.44–0.96"),
    ]
    add_table(doc, ["问题", "根因", "论文粒度修正", "验证结果"], rows, [3.2, 4.3, 5.3, 3.7])

    add_heading(doc, "3.1 PDE 是否写错：已通过 oracle 排除", 2)
    add_bullet(doc, "把 Gkeyll 真值 ∂xq(t,x) 直接输入同一个流体求解器，状态误差约 10⁻⁵、场能 log-RMSE 约 10⁻³；说明三矩方程、符号和空间导数与数据高度一致。")
    add_bullet(doc, "把 dt 从 0.002 改到 0.01、在 Poisson/Ampère 间切换，学习闭合的长期误差几乎不变；因此主要矛盾不是时间步或电场求解器。")
    add_bullet(doc, "FNO 在真值状态上可非常准确，但自身闭环状态逐渐偏离训练流形后，闭合误差被反复反馈并累积成幅值和相位漂移。")

    add_heading(doc, "3.2 为什么最终提升不是简单的参数扫描", 2)
    add_number(doc, "先对齐数据源和非线性案例，保证训练对象确实包含论文要复现的粒子俘获后动力学。")
    add_number(doc, "再对齐中心矩与 ∂xq 标签，消除闭合量本身的语义错误。")
    add_number(doc, "随后用 oracle 验证 PDE，把“流体方程错误”和“学习闭合误差”解耦。")
    add_number(doc, "最后对齐 FNO 结构、RK4 调用粒度和训练帧数，并通过 width/modes 与部署谱截断选择提升闭环稳定性。")
    add_number(doc, "在同轨迹结果达到论文级现象后，再引入严格划分揭示泛化边界，而不是继续用同一条曲线自证。")
    doc.add_page_break()


def add_same_trajectory_results(doc: Document) -> None:
    add_heading(doc, "4. 论文同轨迹复现结果", 1)
    add_callout(
        doc,
        "本章口径",
        "FNO 使用 k=0.35、A=0.10 案例完整 t=0–40 的 8000 帧进行纯监督训练；没有 rollout 微调。随后从相同初值自由推进 20,000 个 RK4 步。"
        "这与论文主非线性案例的展示协议相近，适合回答“FNO 能否学习并嵌入闭合”，但不是严格样本外测试。",
        AMBER,
        NAVY,
    )

    add_heading(doc, "4.1 长时间场能：FNO 复现回升，HP 继续过阻尼", 2)
    add_picture(
        doc,
        IMAGES["field"],
        "图 1  Gkeyll 真值、Fluid+HP 与 Fluid+FNO 的归一化电场能长时间演化。右图中橙线与黑线近乎重合；左图 HP 在非线性阶段出现持续过阻尼和相位偏差。",
        6.9,
        "results/gkeyll_round8/final_report/paper_field_energy_truth_hp_fno.png",
    )
    rows = [
        ("HP", "1.1773", "0.01520", "0.63164", "0.04289", "0"),
        ("纯监督 FNO（最佳 seed）", "0.03791", "0.000186", "0.006253", "0.000755", "0"),
    ]
    add_table(doc, ["闭合", "场能 log10 RMSE", "n rel-L2", "u rel-L2", "p rel-L2", "clamp"], rows, [4.1, 3.0, 2.4, 2.4, 2.4, 1.6])
    p = doc.add_paragraph()
    p.add_run(
        "解释：HP 是线性响应闭合，能给出初期 Landau 阻尼趋势，却缺少粒子俘获造成的非线性能量返还；FNO 从完整动理学轨迹中学习到这种状态依赖的热通量反馈，"
        "因此能恢复场能最低点后的回升和 bounce 振荡。"
    )

    add_heading(doc, "4.2 低阶矩时空结构", 2)
    add_picture(
        doc,
        IMAGES["moments"],
        "图 2  密度 n、速度 u 和压力 p 的时空演化。每一行从左至右依次为 Gkeyll 真值、Fluid+HP、Fluid+FNO。FNO 保持了主要振幅、相位和后期结构，HP 随时间明显衰减。",
        6.55,
        "results/gkeyll_round8/final_report/paper_moments_truth_hp_fno.png",
    )
    add_callout(
        doc,
        "对图 2 的边界说明",
        "这些是低阶矩的流体闭环结果，不是 FNO 生成的相空间分布。三矩模型没有演化 f(x,v,t)，因此不能把 Gkeyll 的相空间涡旋表述为 FNO 直接预测。",
        LIGHT_GRAY,
        NAVY,
    )
    doc.add_page_break()

    add_heading(doc, "4.3 闭合量本身：FNO 学到了什么", 2)
    add_picture(
        doc,
        IMAGES["closure"],
        "图 3  热通量梯度 ∂xq：Gkeyll 真值、HP 闭合和纯监督 FNO 闭合。统计与部署均采用 mode-8 过滤。",
        6.9,
        "results/gkeyll_round8/final_report/paper_closure_truth_hp_fno.png",
    )
    rows = [
        ("HP", "1.1331", "3.6157×10⁻²", "0.4447"),
        ("纯监督 FNO", "0.00458", "1.4607×10⁻⁴", "0.99999"),
    ]
    add_table(doc, ["闭合", "相对 L2", "RMSE", "相关系数"], rows, [4.2, 4.0, 4.0, 4.0])
    p = doc.add_paragraph()
    p.add_run(
        "离线结果说明，在真值低阶矩流形上，FNO 可以近乎逐帧重建 ∂xq；HP 只能给出与真值弱相关的线性近似。"
        "但离线高相关不自动保证闭环泛化，因为闭环输入由模型自己的上一时刻状态产生。"
    )

    add_heading(doc, "4.4 三矩误差的空间—时间分布", 2)
    add_picture(
        doc,
        IMAGES["moment_errors"],
        "图 4  相对 Gkeyll 真值的低阶矩绝对误差。FNO 与 HP 使用相同流体方程和绘图色标；FNO 误差显著更小。",
        5.7,
        "results/gkeyll_round8/final_report/paper_moment_errors_fno_hp.png",
    )
    doc.add_page_break()


def add_stability(doc: Document) -> None:
    add_heading(doc, "5. 稳定性、谱截断与随机种子", 1)
    add_picture(
        doc,
        IMAGES["stability"],
        "图 5  固定 width=128、modes=32 后，不同稳定性损失与随机种子的闭环比较。短时改进并未自动转化为 t=40 的长期稳健性。",
        6.85,
        "results/gkeyll_round8/final_report/stability_constraint_seed_comparison.png",
    )
    add_heading(doc, "5.1 为什么展示图看起来特别好", 2)
    p = doc.add_paragraph()
    p.add_run(
        "图 1–4 使用 Round7 纯监督 width=128、modes=32 的最佳 seed1，场能误差为 0.03791。"
        "同一设置三个 seed 的误差是 0.3102、0.0379、0.1788，中位数为 0.1788。"
        "因此最佳曲线是真实实验结果，但它代表“最佳已实现复现”，不代表随机初始化下的典型表现。"
    )
    rows = [
        ("seed 0", "0.3102", "完成 t=40", "误差较大"),
        ("seed 1", "0.0379", "完成 t=40", "正式复现图使用"),
        ("seed 2", "0.1788", "完成 t=40", "接近三 seed 中位数"),
        ("三 seed 中位数", "0.1788", "—", "更适合作为典型性能"),
    ]
    add_table(doc, ["运行", "场能 log10 RMSE", "数值状态", "解释"], rows, [3.8, 4.0, 3.8, 4.6])

    add_heading(doc, "5.2 部署 mode 8 的意义", 2)
    add_bullet(doc, "网络训练容量为 modes=32，但闭环只输出前 8 个 Fourier 模；这不是论文正文披露的设置，而是本项目通过 8/12/16/24 扫描得到的工程性稳定化。")
    add_bullet(doc, "mode 8 在这个单波数案例上保留主物理模态，同时抑制对 PDE 反馈敏感的高频误差；部署 8 模的场能误差优于 12、16 和 24 模。")
    add_bullet(doc, "它不能被理解为通用规律。对更宽频谱、多尺度或更高维问题，最大模数必须重新由验证集和物理分辨率共同选择。")

    add_heading(doc, "5.3 简单稳定性损失为什么没有继续改善长期结果", 2)
    p = doc.add_paragraph()
    p.add_run(
        "加入 mode-8 部署损失和频谱一致性后，t≤10 的三 seed 极差下降约 85.9%，但 t=40 中位误差反而从 0.1788 增至 0.2559，"
        "且一个 seed 在 t=27.818 失稳。这说明长期失败主要来自自身状态离开监督数据流形后的反馈累积，而不是一步输出频谱不够平滑。"
    )
    doc.add_page_break()


def add_generalization(doc: Document) -> None:
    add_heading(doc, "6. 严格时间外推与跨案例泛化", 1)
    add_callout(
        doc,
        "严格评估结论",
        "同轨迹最好结果证明方法可以复现论文案例；严格划分则表明现有纯监督 FNO 还不是通用闭合。"
        "两者并不矛盾，回答的是不同问题。",
        AMBER,
        NAVY,
    )

    add_heading(doc, "6.1 单案例因果外推：只看 t<24，预测后期非线性", 2)
    add_picture(
        doc,
        IMAGES["temporal"],
        "图 6  单案例时间外推。训练 t=0–24，验证 t=24–30，测试 t=30–40。左：teacher-forced 闭合误差；右：三 seed 自由闭环场能。",
        6.9,
        "results/gkeyll_round9/final_report/single_case_temporal_extrapolation.png",
    )
    rows = [
        ("训练段", "0–24", "闭合误差低", "已见状态"),
        ("验证段", "24–30", "rel-L2 中位 1.6156", "进入未见非线性状态后立即恶化"),
        ("测试段", "30–40", "rel-L2 中位 1.6544", "相关系数中位仅 0.0659"),
        ("闭环", "0–40", "场能误差中位 0.4443", "2/3 完成；1 个 seed 在 37.906 发散"),
    ]
    add_table(doc, ["区间/口径", "时间", "结果", "物理含义"], rows, [3.4, 2.6, 4.8, 5.6])
    p = doc.add_paragraph()
    p.add_run(
        "这说明只靠线性和弱非线性状态无法外推粒子俘获后的闭合分支。论文无需编码历史也能得到好结果，一个关键原因是其 FNO 在训练时已经看过完整 t=0–40 轨迹，"
        "后期状态并非真正未见。"
    )
    doc.add_page_break()

    add_heading(doc, "6.2 整案例留出：训练其他参数，测试未见轨迹", 2)
    add_picture(
        doc,
        IMAGES["cross_case"],
        "图 7  整轨迹留出泛化。左列为未见案例的离线闭合误差；右列为相同案例的长时间闭环场能。",
        6.65,
        "results/gkeyll_round9/final_report/cross_case_generalization.png",
    )
    rows = [
        ("k=.35, A=.075", "0.0967", "0.9956–0.9975", "0.6981", "0.6031–0.7045"),
        ("k=.40, A=.10", "0.0933", "0.9968–0.9977", "0.9625", "0.8803–1.0418"),
    ]
    add_table(doc, ["未见案例", "离线 rel-L2 中位", "相关系数范围", "闭环场能误差中位", "3-seed 范围"], rows, [3.2, 3.5, 3.5, 3.8, 3.1])
    p = doc.add_paragraph()
    p.add_run(
        "跨案例离线相关系数约 0.997，证明完整非线性训练案例具有一定可迁移结构；但约 10% 的瞬时闭合误差被 20,000 步反馈后，"
        "仍形成明显的场能幅值偏高和相位漂移。所有运行都数值稳定且无 clamp，问题重点是闭环动力学不准，而不是求解器爆炸。"
    )

    add_heading(doc, "6.3 评估协议改变了结论的强弱", 2)
    add_picture(
        doc,
        IMAGES["protocol_gap"],
        "图 8  同轨迹完整训练、单案例时间外推与整案例留出的场能误差对比。",
        6.25,
        "results/gkeyll_round9/final_report/protocol_generalization_gap.png",
    )
    rows = [
        ("同轨迹完整训练", "0.1788", "最佳 seed 0.0379", "局部轨迹重建"),
        ("单案例时间外推", "0.4443", "约为最佳同轨迹的 11.7 倍", "后期状态外推失败"),
        ("整案例留出", "约 0.7924", "约为最佳同轨迹的 20.9 倍", "跨参数闭环仍不足"),
    ]
    add_table(doc, ["评估协议", "场能误差中位", "相对差距", "能够支持的结论"], rows, [4.1, 3.5, 4.4, 4.4])
    doc.add_page_break()


def add_conclusions(doc: Document) -> None:
    add_heading(doc, "7. 结论、研究价值与下一步", 1)
    add_heading(doc, "7.1 可以严谨宣称的成果", 2)
    claims = [
        ("可以", "在与论文一致的 Gkeyll 1D1V 非线性单案例上，纯监督 FNO 热通量梯度闭合可在三矩流体方程中复现阻尼、最低点、回升和后期振荡，并显著优于 HP。"),
        ("可以", "PDE、中心矩定义、Ampère 场推进和 RK4 闭合接口已经由真值 closure oracle 验证；当前主要长期误差来自学习闭合的分布漂移。"),
        ("可以", "从 PIC 初试到 Gkeyll 严格复现的改善主要来自物理/数据/数值合同对齐，而不是单纯扩大 FNO。"),
        ("不可以", "不能把最佳同轨迹 seed 表述为通常的样本外性能，也不能据此宣称跨参数通用闭合。"),
        ("不可以", "不能说 FNO 重建了相空间涡旋；流体模型只演化 n、u、p、E，相空间图仅来自 Gkeyll 动理学真值。"),
        ("暂不可以", "不能说已经在 PIC 上达到论文效果。PIC 路线还需要处理粒子噪声、标签平滑和 Vlasov→PIC 域迁移。"),
    ]
    table = add_table(doc, ["口径", "表述"], claims, [2.7, 13.7])
    for row in table.rows[1:]:
        if row.cells[0].text == "可以":
            set_cell_shading(row.cells[0], GREEN)
        else:
            set_cell_shading(row.cells[0], AMBER)

    add_heading(doc, "7.2 对本项目的研究价值", 2)
    add_bullet(doc, "建立了一条可审计的“动理学数据 → 中心矩 → ∂xq → FNO → 三矩流体闭环”完整基线，可作为后续二维/三维闭合研究的最小验证平台。")
    add_bullet(doc, "给出了与经典 HP 在同一求解器内的对照，能够把“非线性闭合价值”从数值离散与 PDE 差异中分离出来。")
    add_bullet(doc, "量化了同轨迹复现与严格泛化之间的性能鸿沟，这本身是面向 AI/计算物理论文的重要结果：离线高相关不等于闭环动力学可靠。")
    add_bullet(doc, "为 PIC 路线提供了正确参照：先在无粒子噪声的 Vlasov 数据上验证物理映射，再研究噪声鲁棒、跨求解器迁移和大规模 CUDA 数据生成。")

    add_heading(doc, "7.3 下一轮优先路线", 2)
    next_rows = [
        ("1", "扩展非线性状态覆盖", "围绕 k=0.30–0.40、A=0.09–0.15 增加完整非线性轨迹，并保持整案例划分", "降低跨参数离线误差"),
        ("2", "闭环感知模型选择", "验证集加入短自由推进、复数主模 Ek 的幅值/相位、场能包络和转折时刻", "避免只按逐帧 MSE 选模型"),
        ("3", "偏移状态训练", "scheduled sampling、离线扰动状态或跨 bounce 周期的多重 shooting", "提高离开真值流形后的恢复能力"),
        ("4", "检验记忆需求", "比较当前状态 FNO 与覆盖至少一个振荡周期的短历史 FNO", "区分状态覆盖不足与闭合非唯一性"),
        ("5", "PIC 域迁移", "使用多 seed 粒子平均、保守平滑/谱截断和 CUDA 生成器，在 case-wise split 下验证", "从 Vlasov 基线迁移到实际 PIC 数据"),
        ("6", "更高维扩展", "在 1D 闭环泛化达标后，再扩展 Gkeyll/PIC 到更高维和更高阶矩", "控制计算量与科学风险"),
    ]
    add_table(doc, ["优先级", "任务", "具体做法", "验收目标"], next_rows, [1.5, 3.5, 7.4, 4.0])

    add_callout(
        doc,
        "最终结论",
        "我们已经把 Huang et al. (2025) 的核心“动理学热通量闭合可驱动流体模型复现非线性 Landau 阻尼”命题，在高度对齐的 Gkeyll 单案例上复现出来；"
        "同时，严格实验表明当前 FNO 的优势仍主要局限于训练状态附近。下一步最有价值的工作不再是继续追求更漂亮的同轨迹曲线，而是让闭合在未见非线性状态和未见参数案例上保持正确的长期能量交换与相位。",
        GREEN,
        TEAL,
    )
    doc.add_page_break()


def add_appendix(doc: Document) -> None:
    add_heading(doc, "附录 A：关键产物与可复核路径", 1)
    rows = [
        ("论文对齐结果摘要", "results/gkeyll_round8/final_report/round8_paper_comparison_summary.json"),
        ("严格泛化结果摘要", "results/gkeyll_round9/final_report/round9_generalization_summary.json"),
        ("Round8 中文报告", "docs/reports/GKEYLL_STABILITY_CONSTRAINT_ROUND8_zh-CN.md"),
        ("Round9 中文报告", "docs/reports/GKEYLL_GENERALIZATION_ROUND9_zh-CN.md"),
        ("PDE 审计报告", "docs/reports/GKEYLL_PDE_CLOSED_LOOP_AUDIT_zh-CN.md"),
        ("严格单案例复现", "docs/reports/GKEYLL_STRICT_REPRODUCTION_ROUND6_zh-CN.md"),
        ("本文档生成脚本", "scripts/create_reproduction_docx.py"),
    ]
    add_table(doc, ["产物", "项目内相对路径"], rows, [5.0, 11.4])

    add_heading(doc, "附录 B：参考文献与复现依据", 1)
    refs = [
        "Huang, Z., Dong, C., & Wang, L. (2025). Machine-learning heat flux closure for multi-moment fluid modeling of nonlinear Landau damping. Proceedings of the National Academy of Sciences, DOI: 10.1073/pnas.2419073122.",
        "Hammett, G. W., & Perkins, F. W. (1990). Fluid moment models for Landau damping with application to the ion-temperature-gradient instability. Physical Review Letters, 64, 3019.",
        "本项目 Round6–Round9 的机器可读 JSON、可视化和中文审计报告。所有数字均由这些本地产物汇总，未从展示图片反向估读。",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(ref)
        set_run_font(r)
        r.font.size = Pt(9.3)

    add_heading(doc, "附录 C：术语", 1)
    terms = [
        ("同轨迹（same trajectory）", "训练中使用了该案例完整时间轨迹；闭环从同一初值重新自由推进。"),
        ("teacher-forced", "每一帧输入都来自真值低阶矩，而不是模型上一时刻预测。"),
        ("closed loop / rollout", "模型输出的闭合进入 PDE；PDE 产生下一状态，再反馈给模型。"),
        ("整案例留出（case-wise split）", "测试参数案例的任何时间帧都不参与训练、归一化统计或 checkpoint 选择。"),
        ("场能 log10 RMSE", "对归一化电场能取 log10 后计算 RMSE，强调跨多个数量级的阻尼和回升差异。"),
        ("clamp", "密度或压力触及数值下限的次数；clamp=0 表示没有依靠钳位维持推进。"),
    ]
    add_table(doc, ["术语", "含义"], terms, [5.0, 11.4])


def build_document() -> Document:
    for path in IMAGES.values():
        if not path.exists():
            raise FileNotFoundError(f"Missing report image: {path}")

    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    configure_styles(doc)
    set_default_language(doc)
    configure_sections(doc)

    properties = doc.core_properties
    properties.title = "Huang et al. (2025) 非线性 Landau 阻尼 FNO 热流闭合复现报告"
    properties.subject = "物理合同对齐、Gkeyll 同轨迹复现、严格时间外推与跨案例泛化"
    properties.author = "landau-damping-surrogate-standardized 项目组"
    properties.keywords = "Landau damping, FNO, Gkeyll, heat flux closure, reproduction"
    properties.comments = "Generated from Round6–Round9 machine-readable results."

    add_title_page(doc)
    add_contents(doc)
    add_executive_summary(doc)
    add_contract(doc)
    add_alignment_journey(doc)
    add_same_trajectory_results(doc)
    add_stability(doc)
    add_generalization(doc)
    add_conclusions(doc)
    add_appendix(doc)
    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
