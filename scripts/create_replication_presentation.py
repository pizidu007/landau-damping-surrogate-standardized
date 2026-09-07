#!/usr/bin/env python3
"""Build the Chinese Huang et al. replication presentation."""
from __future__ import annotations

from pathlib import Path
import json
import math

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/presentations/Huang2025_PIC_FNO_replication_zh-CN.pptx"
ASSET = ROOT / "reports/presentations/assets"
RESULT = ROOT / "results"
FONT = "WenQuanYi Micro Hei"
FONT_MONO = "WenQuanYi Micro Hei Mono"

NAVY = "14213D"
BLUE = "176B87"
CYAN = "27A6C1"
ORANGE = "F28E2B"
RED = "D9534F"
GREEN = "2E8B57"
PURPLE = "7A5195"
INK = "263238"
MUTED = "607D8B"
LIGHT = "F4F7FA"
WHITE = "FFFFFF"
LINE = "DCE4EA"
PALE_BLUE = "E8F4F8"
PALE_ORANGE = "FFF2E3"
PALE_RED = "FCEBEC"
PALE_GREEN = "EAF6EF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_font(run, size: float, color: str = INK, bold: bool = False, font: str = FONT) -> None:
    run.font.name = font
    run._r.get_or_add_rPr().set(qn("a:ea"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.font.bold = bold


def add_text(slide, text: str, x: float, y: float, w: float, h: float, *,
             size: float = 20, color: str = INK, bold: bool = False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin: float = 0.05,
             font: str = FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.clear()
    box.text_frame.margin_left = Inches(margin)
    box.text_frame.margin_right = Inches(margin)
    box.text_frame.margin_top = Inches(margin)
    box.text_frame.margin_bottom = Inches(margin)
    box.text_frame.vertical_anchor = valign
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    set_font(run, size, color, bold, font)
    return box


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float,
                *, size: float = 18, color: str = INK, accent: str = CYAN,
                line_spacing: float = 1.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    for index, item in enumerate(items):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = "•  " + item
        paragraph.space_after = Pt(9)
        paragraph.line_spacing = line_spacing
        run = paragraph.runs[0]
        set_font(run, size, color)
        if index == 0:
            run.font.color.rgb = rgb(color)
    return box


def shape(slide, x: float, y: float, w: float, h: float, *, fill: str = WHITE,
          line: str = LINE, radius=True, transparency: int = 0):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = rgb(fill)
    item.fill.transparency = transparency
    item.line.color.rgb = rgb(line)
    item.line.width = Pt(1)
    return item


def add_card(slide, title: str, body: str, x: float, y: float, w: float, h: float,
             *, fill: str = WHITE, accent: str = CYAN, title_size: float = 16,
             body_size: float = 13, body_color: str = INK):
    shape(slide, x, y, w, h, fill=fill)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(0.07), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = rgb(accent); bar.line.fill.background()
    add_text(slide, title, x + 0.22, y + 0.12, w - 0.32, 0.33,
             size=title_size, bold=True, color=accent)
    add_text(slide, body, x + 0.22, y + 0.52, w - 0.34, h - 0.62,
             size=body_size, color=body_color)


def add_metric(slide, value: str, label: str, x: float, y: float, w: float,
               *, color: str = CYAN, fill: str = WHITE):
    shape(slide, x, y, w, 1.03, fill=fill)
    add_text(slide, value, x + 0.1, y + 0.10, w - 0.2, 0.44,
             size=25, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.1, y + 0.59, w - 0.2, 0.27,
             size=11.5, color=MUTED, align=PP_ALIGN.CENTER)


def add_title(slide, title: str, subtitle: str | None = None, *, section: str = "本项目结果"):
    add_text(slide, title, 0.65, 0.33, 11.4, 0.48, size=25, bold=True, color=NAVY)
    if subtitle:
        add_text(slide, subtitle, 0.68, 0.82, 11.4, 0.30, size=11.5, color=MUTED)
    tag_fill = PURPLE if section == "论文原图" else (ORANGE if section == "诊断" else BLUE)
    pill = shape(slide, 11.70, 0.36, 0.98, 0.33, fill=tag_fill, line=tag_fill)
    add_text(slide, section, 11.72, 0.40, 0.94, 0.20, size=9.5, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                  Inches(0.65), Inches(1.11), Inches(12.03), Inches(0.025))
    line.fill.solid(); line.fill.fore_color.rgb = rgb(LINE); line.line.fill.background()


def add_footer(slide, number: int, source: str = ""):
    add_text(slide, source, 0.68, 7.12, 10.8, 0.20, size=8.5, color=MUTED)
    add_text(slide, f"{number:02d}", 12.05, 7.09, 0.55, 0.22, size=9.5,
             color=MUTED, align=PP_ALIGN.RIGHT)


def image_dimensions(path: Path) -> tuple[int, int]:
    with Image.open(path) as image:
        return image.size


def add_image(slide, path: Path, x: float, y: float, w: float, h: float,
              *, card: bool = True, pad: float = 0.08, caption: str | None = None):
    if not path.exists():
        raise FileNotFoundError(path)
    if card:
        shape(slide, x, y, w, h, fill=WHITE)
    width, height = image_dimensions(path)
    cap_h = 0.28 if caption else 0.0
    area_w, area_h = w - 2 * pad, h - 2 * pad - cap_h
    scale = min(area_w / width, area_h / height)
    draw_w, draw_h = width * scale, height * scale
    left = x + (w - draw_w) / 2
    top = y + pad + (area_h - draw_h) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(draw_w), Inches(draw_h))
    if caption:
        add_text(slide, caption, x + 0.12, y + h - 0.27, w - 0.24, 0.18,
                 size=9.5, color=MUTED, align=PP_ALIGN.CENTER)


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, *, color: str = CYAN, width: float = 2.2):
    arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                       Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    arrow.line.color.rgb = rgb(color)
    arrow.line.width = Pt(width)
    arrow.line.end_arrowhead = True
    return arrow


def crop_asset(source: Path, target: Path, box_fraction: tuple[float, float, float, float]) -> None:
    with Image.open(source) as image:
        w, h = image.size
        box = tuple(int(value * dim) for value, dim in zip(box_fraction, (w, h, w, h)))
        image.crop(box).save(target)


def contact_sheet(paths: list[Path], labels: list[str], target: Path,
                  columns: int = 3, size: tuple[int, int] = (1800, 1000)) -> None:
    canvas = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(canvas)
    try:
        font = ImageFont.truetype("/usr/share/fonts/wqy-microhei/wqy-microhei.ttc", 24)
    except OSError:
        font = ImageFont.load_default()
    rows = math.ceil(len(paths) / columns)
    cell_w, cell_h = size[0] // columns, size[1] // rows
    for index, (path, label) in enumerate(zip(paths, labels)):
        row, column = divmod(index, columns)
        with Image.open(path) as source:
            image = source.convert("RGB")
            available = (cell_w - 24, cell_h - 48)
            image.thumbnail(available, Image.Resampling.LANCZOS)
            left = column * cell_w + (cell_w - image.width) // 2
            top = row * cell_h + 36 + (cell_h - 42 - image.height) // 2
            canvas.paste(image, (left, top))
        draw.text((column * cell_w + 12, row * cell_h + 6), label, font=font, fill="#263238")
    canvas.save(target, quality=92)


def new_slide(prs: Presentation, background: str = LIGHT):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = rgb(background)
    return slide


def add_table(slide, rows: list[list[str]], x: float, y: float, w: float, h: float,
              widths: list[float] | None = None, header_fill: str = NAVY):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if widths:
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(header_fill if r == 0 else (WHITE if r % 2 else "EEF3F6"))
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
                for run in paragraph.runs:
                    set_font(run, 11.5 if r else 12.0, WHITE if r == 0 else INK, r == 0)
    return table_shape


def build_assets() -> None:
    ASSET.mkdir(parents=True, exist_ok=True)
    crop_asset(ASSET / "paper_page-06.png", ASSET / "paper_fig1_fno.png", (0.10, 0.09, 0.90, 0.81))
    crop_asset(ASSET / "paper_page-07.png", ASSET / "paper_fig2_closure.png", (0.14, 0.035, 0.87, 0.82))
    crop_asset(ASSET / "paper_page-08.png", ASSET / "paper_fig3_field.png", (0.09, 0.28, 0.91, 0.76))
    crop_asset(ASSET / "paper_page-09.png", ASSET / "paper_fig4_phase.png", (0.10, 0.05, 0.90, 0.86))

    v1_train = sorted((RESULT / "closure_fno_v1/stage3_supervised").glob("*/training_curve.png"))
    contact_sheet(v1_train, [path.parent.name for path in v1_train], ASSET / "v1_training_contact.png", columns=4)
    history = sorted((RESULT / "closure_fno_v2/stage4_filtered_history").glob("*/training_curve.png"))
    contact_sheet(history, [path.parent.name for path in history], ASSET / "history_training_contact.png", columns=3, size=(1800, 620))
    v3 = sorted((RESULT / "closure_fno_v3_raw_single/stage1_training").glob("seed*/closure_xt.png"))
    contact_sheet(v3, [path.parent.name for path in v3], ASSET / "v3_seed_contact.png", columns=3, size=(1800, 650))


def build_presentation() -> Path:
    build_assets()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slides: list = []

    def finish(slide, source: str = ""):
        slides.append(slide)
        add_footer(slide, len(slides), source)

    # 1 Cover
    slide = new_slide(prs, NAVY)
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    band.fill.solid(); band.fill.fore_color.rgb = rgb(ORANGE); band.line.fill.background()
    add_text(slide, "Huang 等（2025）\nFNO 热流闭合复现汇报", 0.75, 1.10, 7.1, 1.55,
             size=31, bold=True, color=WHITE)
    add_text(slide, "从 Vlasov 论文路线到 CUDA-PIC 数据：\n离线闭合、非线性相空间与长时自由演化", 0.80, 2.92, 6.75, 1.05,
             size=18, color="C9D8E6")
    add_card(slide, "一句话结论", "瞬时 ∂xM₃ 拟合可达 4.8% test 误差；\n纯 FNO 闭环仍未复现论文的长期稳定性。",
             0.80, 4.44, 6.1, 1.20, fill="20304D", accent=ORANGE, title_size=14,
             body_size=14, body_color=WHITE)
    add_image(slide, ASSET / "paper_cover.png", 8.25, 0.48, 4.15, 6.30, card=True)
    add_text(slide, "2026-08-28", 0.82, 6.78, 2.2, 0.25, size=10.5, color="9FB3C8")
    finish(slide, "论文：Huang, Dong & Wang, PNAS (2025), DOI 10.1073/pnas.2419073122")

    # 2 Executive summary
    slide = new_slide(prs); add_title(slide, "结论先行：复现了闭合映射，没有复现长期闭环")
    add_card(slide, "成功 ①", "公开 Gkeyll 数据同轨迹交错 test\n相对 L2 = 0.00279", 0.75, 1.45, 2.85, 1.30, fill=PALE_GREEN, accent=GREEN)
    add_card(slide, "成功 ②", "PIC 原始矩单帧 FNO\ntest L2 = 0.04798，corr = 0.99885", 3.78, 1.45, 2.85, 1.30, fill=PALE_GREEN, accent=GREEN)
    add_card(slide, "物理证据", "k=0.35, A=0.10 在 t=30–60\n出现俘获、卷曲与细丝化", 6.81, 1.45, 2.85, 1.30, fill=PALE_BLUE, accent=BLUE)
    add_card(slide, "未复现", "标准 mode-24 纯 FNO\n自由演化仅到 t=1.8", 9.84, 1.45, 2.75, 1.30, fill=PALE_RED, accent=RED)
    add_bullets(slide, [
        "论文证明的是 Vlasov 数据上的 Fluid + ML；我们检验的是更有噪声的 PIC 矩闭合，两者物理目标一致、离散数据分布不同。",
        "历史 5 帧将 PIC 离线误差从 0.234 降到 0.173，但离线提升没有自动变成长时稳定。",
        "Ampère/Poisson、dt=0.01/0.005 均未改变 mode-24 的失稳时间，主瓶颈不是时间步。",
        "真实 PIC ∂xM₃ 的 oracle 闭合也只能到 t=8.7：粗粒化矩方程与 PIC 轨迹并非严格自治系统。",
    ], 0.90, 3.25, 11.65, 2.65, size=17)
    add_card(slide, "当前判断", "这不是“FNO 没学会热通量”，而是“监督误差没有约束离流形闭环稳定性”。",
             1.45, 6.15, 10.45, 0.62, fill="EAF0F5", accent=ORANGE, title_size=13, body_size=13)
    finish(slide)

    # 3 Paper method
    slide = new_slide(prs); add_title(slide, "论文路线：用 FNO 把低阶矩映射到热流梯度", section="论文原图")
    add_image(slide, ASSET / "paper_fig1_fno.png", 0.70, 1.35, 7.45, 5.45, card=True)
    add_card(slide, "输入", "当前时刻整段空间场\nn(x), u(x), p(x)", 8.45, 1.55, 3.95, 1.05, fill=PALE_BLUE, accent=BLUE)
    add_card(slide, "算子", "P → 4 个 Fourier layer → Q\n同时学习全局谱信息与局部变换", 8.45, 2.84, 3.95, 1.22, fill="F2ECF7", accent=PURPLE)
    add_card(slide, "输出", "∂q/∂x 作为压力方程源项\n每个 RK4 时间步重新调用 FNO", 8.45, 4.30, 3.95, 1.22, fill=PALE_ORANGE, accent=ORANGE)
    add_card(slide, "关键特征", "无显式历史；非局域；闭环使用新预测状态。", 8.45, 5.76, 3.95, 0.86, fill=PALE_GREEN, accent=GREEN)
    finish(slide, "论文 Fig. 1；输入/输出说明据 Huang et al. (2025)")

    # 4 paper numerical setup and target closure
    slide = new_slide(prs); add_title(slide, "论文基准：同一条非线性 Vlasov 轨迹上的强插值能力", section="论文原图")
    add_image(slide, ASSET / "paper_fig2_closure.png", 0.70, 1.35, 7.35, 5.55, card=True)
    add_metric(slide, "k=0.35", "基波波数", 8.43, 1.55, 1.83, color=BLUE, fill=PALE_BLUE)
    add_metric(slide, "A=0.10", "非线性扰动", 10.52, 1.55, 1.83, color=ORANGE, fill=PALE_ORANGE)
    add_metric(slide, "64×64", "x–v 网格", 8.43, 2.85, 1.83, color=PURPLE, fill="F2ECF7")
    add_metric(slide, "8000", "训练快照，Δt=0.005", 10.52, 2.85, 1.83, color=GREEN, fill=PALE_GREEN)
    add_bullets(slide, [
        "GKEYLL Vlasov–Ampère 数据覆盖 t=0–40。",
        "论文报告 ∂q/∂x 绝对误差约 10⁻⁴。",
        "Fluid + ML 使用更小的 Δt=0.002 外推到 20000 步。",
        "训练、验证和测试都来自同一条轨迹。",
    ], 8.48, 4.30, 3.85, 2.05, size=14.5)
    finish(slide, "论文 Fig. 2；参数与时间步来自正文")

    # 5 paper target dynamics
    slide = new_slide(prs); add_title(slide, "论文目标结果：场能再增长与非线性相空间洞", section="论文原图")
    add_image(slide, ASSET / "paper_fig3_field.png", 0.68, 1.35, 5.95, 4.90, card=True, caption="论文 Fig. 3：Fluid + ML 几乎重合 Vlasov 场能")
    add_image(slide, ASSET / "paper_fig4_phase.png", 6.83, 1.35, 5.82, 4.90, card=True, caption="论文 Fig. 4：非线性面板只画正共振速度窗")
    add_card(slide, "我们真正要复现的不是一张热图", "接受标准必须同时包含：∂q/∂x 离线精度、t=60 场能相位/包络、低阶矩稳定性与非线性相空间证据。",
             1.15, 6.42, 11.05, 0.55, fill=PALE_ORANGE, accent=ORANGE, title_size=12.5, body_size=12.5)
    finish(slide, "论文 Fig. 3–4")

    # 6 Scope diagram
    slide = new_slide(prs); add_title(slide, "我们的复现：物理问题相同，动理学数据源不同")
    add_card(slide, "论文", "GKEYLL\nEulerian Vlasov\n低噪声 f(x,v)", 0.90, 1.65, 2.55, 1.55, fill="F2ECF7", accent=PURPLE, title_size=15, body_size=17)
    add_card(slide, "本项目", "CUDA-PIC\n26,214,400 particles/case\n3 seeds 平均", 0.90, 4.15, 2.55, 1.55, fill=PALE_BLUE, accent=BLUE, title_size=15, body_size=16)
    add_arrow(slide, 3.55, 2.40, 5.10, 2.40, color=PURPLE)
    add_arrow(slide, 3.55, 4.90, 5.10, 4.90, color=BLUE)
    add_card(slide, "共同低阶状态", "n, u, p 或原始矩 M₀, M₁, M₂", 5.15, 2.73, 2.80, 1.35, fill=WHITE, accent=CYAN, title_size=15, body_size=16)
    add_arrow(slide, 8.05, 3.40, 9.12, 3.40, color=CYAN)
    add_card(slide, "FNO closure", "预测 q 或 ∂xM₃\n嵌入三矩流体方程", 9.18, 2.63, 2.85, 1.55, fill=PALE_ORANGE, accent=ORANGE, title_size=15, body_size=16)
    add_card(slide, "PIC 带来的额外难点", "有限粒子噪声随矩阶和 Fourier mode 增长；seed 平均不能消除离散乘积偏差；f→低阶矩是多对一映射。",
             4.15, 5.15, 7.85, 1.10, fill=PALE_RED, accent=RED, title_size=14, body_size=13.5)
    add_text(slide, "因此：我们是在 PIC 上检验同一个闭合思想，但不是对论文数值离散的逐字节复刻。",
             1.15, 6.55, 11.0, 0.36, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    finish(slide)

    # 7 roadmap
    slide = new_slide(prs); add_title(slide, "三轮实验路线：从可用闭合到严格论文式复现")
    steps = [
        ("v1", "PIC 工程基线", "20 参数对 / 60 cases\n预测中心热流 q\n75% HP 混合稳定到 t=60", BLUE),
        ("v2", "论文与历史审计", "公开 Gkeyll 数据复核\nraw/central moment 纠偏\nhistory-5 离线提升", PURPLE),
        ("v3", "单轨迹原始矩", "k=.35, A=.10\n单帧直接 ∂xM₃\n自由闭环谱消融", ORANGE),
    ]
    for i, (tag, title, body, color) in enumerate(steps):
        x = 0.95 + i * 4.05
        add_card(slide, f"{tag} · {title}", body, x, 2.0, 3.30, 2.25,
                 fill=WHITE, accent=color, title_size=17, body_size=15)
        add_metric(slide, ["工程可用", "问题定位", "严格否证"][i], "该轮主要产出", x + 0.42, 4.58, 2.46, color=color,
                   fill=[PALE_BLUE, "F2ECF7", PALE_ORANGE][i])
        if i < 2:
            add_arrow(slide, x + 3.38, 3.11, x + 4.00, 3.11, color=MUTED)
    add_card(slide, "方法论变化", "指标从“闭合场看起来像”升级为“同轨迹插值 + 因果外推 + 自由闭环 + 物理守恒 + 相空间证据”。",
             1.20, 6.08, 10.93, 0.64, fill="EAF0F5", accent=CYAN, title_size=13, body_size=13)
    finish(slide)

    # 8 dataset
    slide = new_slide(prs); add_title(slide, "CUDA-PIC 数据集：为非线性闭合准备高粒子数标签")
    add_metric(slide, "60", "正式 PIC cases", 0.85, 1.42, 2.0, color=BLUE, fill=PALE_BLUE)
    add_metric(slide, "26.2M", "particles / case", 3.05, 1.42, 2.0, color=GREEN, fill=PALE_GREEN)
    add_metric(slide, "Nx=128", "矩网格", 5.25, 1.42, 2.0, color=PURPLE, fill="F2ECF7")
    add_metric(slide, "t=60", "长时间终点", 7.45, 1.42, 2.0, color=ORANGE, fill=PALE_ORANGE)
    add_metric(slide, "3", "seeds / 参数", 9.65, 1.42, 2.0, color=CYAN, fill=PALE_BLUE)
    add_bullets(slide, [
        "每个样本同时保存原始矩 M₀–M₃、中心矩 n/u/p/q、电场能与相空间帧。",
        "生成器强制 CUDA；moment Δt=0.1、phase-space Δt=1.0。",
        "相空间使用 Nx=512、速度轴 1537 点，便于观察细丝化与粒子俘获。",
        "正式数据集中已有 k=0.35, A=0.10 的三 seed 非线性长轨迹。",
    ], 0.95, 3.08, 5.90, 2.20, size=16)
    add_card(slide, "数据合同", "训练数据、模型和图均位于统一 dataset/project 目录；HDF5 写入 COMPLETE 状态、参数 JSON 和网格元数据。",
             0.95, 5.54, 5.85, 0.92, fill=WHITE, accent=GREEN, title_size=14, body_size=12.5)
    add_image(slide, RESULT / "closure_fno_v1/stage1_label_audit/plots/seed_noise_parameter_map.png",
              7.15, 3.02, 5.30, 3.45, caption="不同参数下 PIC seed 噪声/信号")
    finish(slide)

    # 9 label audit
    slide = new_slide(prs); add_title(slide, "标签审计：高阶矩的高频部分最容易被 PIC 噪声污染")
    add_image(slide, RESULT / "closure_fno_v1/stage1_label_audit/plots/dqdx_signal_noise_spectrum.png",
              0.72, 1.37, 6.10, 4.75, caption="∂xq 的信号与 seed 噪声随 Fourier mode 的分布")
    add_image(slide, RESULT / "closure_fno_v1/stage1_label_audit/plots/seed_noise_parameter_map.png",
              6.98, 1.37, 5.65, 4.75, caption="参数图：强弱扰动下的 seed 噪声差异")
    add_metric(slide, "1.289", "v1 median noise/signal", 1.35, 6.28, 2.55, color=RED, fill=PALE_RED)
    add_metric(slide, "mode 23", "审计推荐最大模式", 4.15, 6.28, 2.55, color=ORANGE, fill=PALE_ORANGE)
    add_card(slide, "训练策略", "三 seed 平均 + 谱损失 + mode-24 截断；避免让网络拟合粒子噪声。",
             7.08, 6.28, 4.75, 1.03, fill=PALE_GREEN, accent=GREEN, title_size=13, body_size=12.5)
    finish(slide)

    # 10 v1 offline closure
    slide = new_slide(prs); add_title(slide, "v1：PIC 热流闭合可以显著优于传统基线")
    add_image(slide, RESULT / "closure_fno_v1/final_visualizations/group1_closure/dqdx_truth_prediction_error_xt.png",
              0.65, 1.35, 8.20, 4.95, caption="PIC truth / FNO closure / absolute error（横轴时间，纵轴空间）")
    add_image(slide, RESULT / "closure_fno_v1/final_visualizations/group1_closure/dqdx_lineouts.png",
              9.02, 1.35, 3.65, 3.00, caption="典型时刻空间切片")
    add_metric(slide, "0.309", "pure FNO test L2", 9.15, 4.65, 1.53, color=BLUE, fill=PALE_BLUE)
    add_metric(slide, "0.985", "HP test L2", 10.91, 4.65, 1.53, color=RED, fill=PALE_RED)
    add_card(slide, "离线结论", "纯 FNO 相对 HP 改善 68.7%，相关系数 0.952。",
             9.15, 5.88, 3.30, 0.72, fill=PALE_GREEN, accent=GREEN, title_size=13, body_size=12)
    finish(slide)

    # 11 v1 dynamics
    slide = new_slide(prs); add_title(slide, "v1：低阶矩与场能闭环——混合方案可稳定到 t=60")
    add_image(slide, RESULT / "closure_fno_v1/final_visualizations/group2_dynamics/electric_energy_rollout.png",
              0.70, 1.35, 5.40, 4.85, caption="电场能：FNO/HP 混合、HP、zero 与 PIC")
    add_image(slide, RESULT / "closure_fno_v1/final_visualizations/group2_dynamics/fluid_moments_xt.png",
              6.28, 1.35, 6.35, 4.85, caption="n/u/p 的时空演化对比")
    add_metric(slide, "t=60", "5/5 test rollouts 完成", 1.15, 6.35, 2.65, color=GREEN, fill=PALE_GREEN)
    add_metric(slide, "1.308", "场能 log10 RMSE", 4.05, 6.35, 2.65, color=ORANGE, fill=PALE_ORANGE)
    add_card(slide, "部署选择", "75% HP + 25% FNO，mode-24，流体子步 dt=0.02。",
             7.18, 6.35, 4.85, 1.03, fill=PALE_BLUE, accent=BLUE, title_size=13, body_size=12.5)
    finish(slide)

    # 12 stability
    slide = new_slide(prs); add_title(slide, "v1：稳定不等于准确——长期误差仍是主要缺口", section="诊断")
    add_image(slide, RESULT / "closure_fno_v1/final_visualizations/group3_stability/rollout_error_by_parameter.png",
              0.75, 1.38, 5.85, 4.55, caption="参数分组的 rollout 场能误差")
    add_image(slide, RESULT / "closure_fno_v1/final_visualizations/group3_stability/total_energy_drift.png",
              6.78, 1.38, 5.80, 4.55, caption="总能量相对漂移")
    add_card(slide, "通过", "所有 test rollout 到 t=60；正性与总能量漂移受控。",
             1.05, 6.12, 3.45, 0.80, fill=PALE_GREEN, accent=GREEN, title_size=13, body_size=12)
    add_card(slide, "未通过", "相对 HP 的场能改善只有 19.6%，未达到 50% 接受门槛。",
             4.92, 6.12, 3.45, 0.80, fill=PALE_RED, accent=RED, title_size=13, body_size=12)
    add_card(slide, "解释", "HP 混合提供耗散稳定性，但会压制纯 FNO 想学习的非线性动力学。",
             8.78, 6.12, 3.45, 0.80, fill=PALE_ORANGE, accent=ORANGE, title_size=13, body_size=12)
    finish(slide)

    # 13 public data
    slide = new_slide(prs); add_title(slide, "v2：公开 Gkeyll 数据能复现“漂亮热图”，但评价协议决定结论")
    add_image(slide, RESULT / "closure_fno_v2/stage1_huang_reproduction/paper_like_dt0005_seed0/closure_xt.png",
              0.65, 1.35, 6.02, 4.85, caption="论文式交错抽帧：同轨迹插值")
    add_image(slide, RESULT / "closure_fno_v2/stage1_huang_reproduction/causal_dt0005_seed0/closure_xt.png",
              6.78, 1.35, 5.90, 4.85, caption="因果划分：t=30–40 非线性时间外推")
    add_metric(slide, "0.00279", "paper-like test L2", 1.55, 6.35, 2.72, color=GREEN, fill=PALE_GREEN)
    add_metric(slide, "0.338", "causal nonlinear test L2", 5.30, 6.35, 2.72, color=RED, fill=PALE_RED)
    add_card(slide, "结论", "同轨迹交错 test 衡量插值能力，不能替代非线性时间外推或自由闭环。",
             8.48, 6.35, 3.75, 1.03, fill=PALE_ORANGE, accent=ORANGE, title_size=13, body_size=12)
    finish(slide)

    # 14 semantics audit
    slide = new_slide(prs); add_title(slide, "v2：变量语义纠偏——公开 p_new / q_new 实际是原始矩", section="诊断")
    rows = [
        ["公开变量", "最初直觉", "方程审计后的解释", "t=0–24 raw 方程残差"],
        ["n_new", "密度 n", "M₀ = ∫f dv", "M₀: 0.0419"],
        ["u_new", "速度 u", "u = M₁/M₀", "M₁: 0.0337"],
        ["p_new", "中心压力 p", "M₂ = ∫v²f dv", "M₂: 0.0297"],
        ["q_new", "中心热流 q", "M₃ = ∫v³f dv", "用于闭合"],
    ]
    add_table(slide, rows, 0.78, 1.50, 11.80, 2.65, widths=[1.8, 2.25, 4.25, 3.5])
    add_card(slide, "转换关系", "p = M₂ − n u²\nq = M₃ − 3uM₂ + 2nu³", 0.95, 4.55, 3.55, 1.22, fill=PALE_BLUE, accent=BLUE, title_size=15, body_size=17)
    add_card(slide, "为什么重要", "变量定义错一处，FNO 离线图仍可能很好，但嵌入的流体方程会完全不自洽。",
             4.78, 4.55, 3.55, 1.22, fill=PALE_RED, accent=RED, title_size=15, body_size=14)
    add_card(slide, "纠偏后仍未解决", "公开数据 raw-moment 闭环 mode-8 只到 t=12.225，说明还有未公开的数值耗散/滤波细节。",
             8.61, 4.55, 3.55, 1.22, fill=PALE_ORANGE, accent=ORANGE, title_size=15, body_size=14)
    add_text(slide, "这一步把“实现错误”与“模型/离散本身不稳定”分开。", 1.0, 6.25, 11.3, 0.38,
             size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    finish(slide)

    # 15 PIC Gkeyll comparison
    slide = new_slide(prs); add_title(slide, "v2：PIC 与 Gkeyll 在线性期一致，非线性高阶矩逐渐分离")
    add_image(slide, RESULT / "closure_fno_v2/stage3_pic_gkeyll_audit/pic_gkeyll_dqdx_xt.png",
              0.70, 1.35, 8.15, 5.55, caption="时间/网格匹配的三 seed CUDA-PIC 与公开 Gkeyll ∂xq")
    add_metric(slide, "0.972", "线性期 ∂xq corr", 9.15, 1.62, 2.90, color=GREEN, fill=PALE_GREEN)
    add_metric(slide, "0.726", "非线性期 ∂xq corr", 9.15, 2.95, 2.90, color=ORANGE, fill=PALE_ORANGE)
    add_metric(slide, "0.331", "mode-24 seed noise/signal", 9.15, 4.28, 2.90, color=RED, fill=PALE_RED)
    add_card(slide, "物理含义", "PIC 不是给 Vlasov 标签“换个求解器”：有限粒子噪声和数值耗散会改变高阶矩。",
             9.05, 5.72, 3.15, 0.94, fill=WHITE, accent=PURPLE, title_size=13, body_size=12)
    finish(slide)

    # 16 history
    slide = new_slide(prs); add_title(slide, "v2：历史信息改善 PIC 离线闭合，但仍不足以保证闭环")
    add_image(slide, RESULT / "closure_fno_v2/final_visualizations/closure_v2_summary.png",
              0.68, 1.35, 6.42, 5.15, caption="第二轮汇总：公开数据协议、PIC 历史消融、长期混合闭环")
    add_image(slide, ASSET / "history_training_contact.png", 7.22, 1.35, 5.43, 2.80,
              caption="单帧 / history-3 / history-5 训练曲线")
    add_metric(slide, "0.234", "单帧 validation L2", 7.42, 4.45, 1.48, color=RED, fill=PALE_RED)
    add_metric(slide, "0.183", "history-3", 9.04, 4.45, 1.48, color=ORANGE, fill=PALE_ORANGE)
    add_metric(slide, "0.173", "history-5", 10.66, 4.45, 1.48, color=GREEN, fill=PALE_GREEN)
    add_card(slide, "解释", "低阶矩到热流并非严格马尔可夫映射；历史帮助消歧。但自由 history-FNO 仍在 t≈24–34 失稳。",
             7.35, 5.75, 4.65, 0.92, fill=WHITE, accent=PURPLE, title_size=13, body_size=12)
    finish(slide)

    # 17 pure long
    slide = new_slide(prs); add_title(slide, "v2：纯 FNO 长时场能——早期贴合不等于长期可信", section="诊断")
    add_image(slide, RESULT / "closure_fno_v2/pure_fno_longtime/pure_fno_field_energy_t60.png",
              0.70, 1.35, 7.25, 5.50, caption="多案例 PIC 与纯 history-FNO 场能演化")
    add_image(slide, RESULT / "closure_fno_v2/pure_fno_longtime/k0p350_a0p100_pure/k0p350_a0p100_field_energy.png",
              8.10, 1.35, 4.55, 3.75, caption="强案例 k=0.35, A=0.10")
    add_card(slide, "观察", "模型能够短期跟踪阻尼振荡，但相位误差积累后包络偏离，最终出现正性/高频失稳。",
             8.20, 5.42, 4.20, 1.08, fill=PALE_RED, accent=RED, title_size=14, body_size=13)
    finish(slide)

    # 18 nonlinearity audit
    slide = new_slide(prs); add_title(slide, "非线性证据：完整速度范围会把相空间结构“淹没”")
    add_image(slide, RESULT / "closure_fno_v2/nonlinearity_audit/k0p350_a0p100_s00/phase_space_full_vs_delta.png",
              0.66, 1.35, 8.05, 5.65, caption="上：完整 f；下：f−〈f〉ₓ。背景扣除后细丝化清晰可见")
    gif = RESULT / "closure_fno_v2/pure_fno_longtime/k0p400_a0p050_pic_vs_pure_fno_maxwellian.gif"
    add_image(slide, gif, 8.90, 1.35, 3.75, 3.35, caption="放映时播放：PIC 与矩匹配 Maxwellian")
    add_card(slide, "绘图原则", "完整 v∈[−6,6] 主要显示 Maxwellian 背景；研究粒子俘获应画 δf 或裁剪到共振速度窗。",
             8.92, 5.02, 3.70, 1.26, fill=PALE_ORANGE, accent=ORANGE, title_size=14, body_size=13)
    finish(slide)

    # 19 resonant phase
    slide = new_slide(prs); add_title(slide, "正共振速度窗：我们的 PIC 确实进入非线性阶段")
    add_image(slide, RESULT / "closure_fno_v2/nonlinearity_audit/k0p350_a0p100_s00/phase_space_resonant_velocity_zoom.png",
              0.66, 1.35, 11.98, 5.55, caption="k=0.35, A=0.10；v=2–4.5；t=20,30,40,50,60")
    add_card(slide, "回答此前疑问", "t≈30 后可见波粒俘获、卷曲与交叉细丝；此前“只有条纹”主要是绘图速度范围过宽。",
             1.15, 6.38, 10.85, 0.61, fill=PALE_GREEN, accent=GREEN, title_size=13, body_size=13)
    finish(slide)

    # 20 v3 setup
    slide = new_slide(prs); add_title(slide, "v3：最接近论文的单轨迹、单帧、原始矩实验")
    add_card(slide, "数据", "k=.35, A=.10\n3 PIC seeds 平均\n601 帧，t=0–60", 0.80, 1.55, 2.42, 1.78, fill=PALE_BLUE, accent=BLUE, title_size=15, body_size=16)
    add_arrow(slide, 3.28, 2.43, 4.05, 2.43, color=BLUE)
    add_card(slide, "输入", "单帧 [M₀, u, M₂]\n无 k、无历史", 4.10, 1.55, 2.42, 1.78, fill="F2ECF7", accent=PURPLE, title_size=15, body_size=17)
    add_arrow(slide, 6.58, 2.43, 7.35, 2.43, color=PURPLE)
    add_card(slide, "FNO", "width 64\n24 modes / 4 layers", 7.40, 1.55, 2.42, 1.78, fill=PALE_ORANGE, accent=ORANGE, title_size=15, body_size=17)
    add_arrow(slide, 9.88, 2.43, 10.65, 2.43, color=ORANGE)
    add_card(slide, "输出", "直接预测 mode-24\n∂xM₃", 10.70, 1.55, 1.88, 1.78, fill=PALE_GREEN, accent=GREEN, title_size=15, body_size=17)
    add_metric(slide, "4.41%", "validation relative L2", 1.05, 4.05, 2.30, color=GREEN, fill=PALE_GREEN)
    add_metric(slide, "4.80%", "test relative L2", 3.70, 4.05, 2.30, color=GREEN, fill=PALE_GREEN)
    add_metric(slide, "0.99885", "test correlation", 6.35, 4.05, 2.30, color=BLUE, fill=PALE_BLUE)
    add_metric(slide, "seed 1", "三随机种子中最佳", 9.00, 4.05, 2.30, color=PURPLE, fill="F2ECF7")
    add_card(slide, "评价协议", "同一轨迹按时间索引 modulo 10 做 8/1/1 交错划分，因此仍属于插值试验。",
             1.30, 5.58, 10.55, 0.90, fill=WHITE, accent=CYAN, title_size=14, body_size=13)
    finish(slide)

    # 21 v3 offline
    slide = new_slide(prs); add_title(slide, "v3 离线结果：PIC truth 与单帧 FNO 几乎重合")
    add_image(slide, RESULT / "closure_fno_v3_raw_single/stage1_training/seed1/closure_xt.png",
              0.65, 1.35, 11.95, 5.70, caption="左：PIC ∂xM₃；中：memoryless FNO；右：绝对误差")
    add_card(slide, "但这个图回答的只是", "“真实轨迹上的当前低阶矩能否插值出当前 ∂xM₃？”——尚未回答自由运行时的稳定性。",
             1.10, 6.35, 11.05, 0.60, fill=PALE_ORANGE, accent=ORANGE, title_size=12.5, body_size=12.5)
    finish(slide)

    # 22 long field
    slide = new_slide(prs); add_title(slide, "v3 长时场能：标准论文式闭环在 t=1.8 失稳", section="诊断")
    add_image(slide, RESULT / "closure_fno_v3_raw_single/final_plots/long_time_field_energy.png",
              0.65, 1.35, 9.15, 5.65, caption="竖虚线为各 FNO 方案停止时间；黑线 PIC 延伸至 t=60")
    add_metric(slide, "1.8", "mode-24 停止时间", 10.12, 1.60, 2.02, color=RED, fill=PALE_RED)
    add_metric(slide, "15.5", "mode-12 停止时间", 10.12, 2.90, 2.02, color=ORANGE, fill=PALE_ORANGE)
    add_metric(slide, "34.5", "mode-8 停止时间", 10.12, 4.20, 2.02, color=ORANGE, fill=PALE_ORANGE)
    add_card(slide, "mode-8 也不成功", "总能量相对跨度 28.62；延长时间来自强低通，并非正确物理。",
             9.98, 5.55, 2.32, 1.10, fill=PALE_RED, accent=RED, title_size=12.5, body_size=11)
    finish(slide)

    # 23 phase result
    slide = new_slide(prs); add_title(slide, "v3 PIC 相空间：固定 v=2–4.5 后，非线性涡旋非常清楚")
    add_image(slide, RESULT / "closure_fno_v3_raw_single/final_plots/pic_phase_space_resonant_velocity.png",
              0.65, 1.35, 12.00, 5.65, caption="t=30 / 40 / 50 / 60；同一色标、同一正共振速度窗")
    add_card(slide, "物理解读", "粒子俘获岛持续卷曲并产生细丝；场能在 t≈23 后再增长，与非线性波粒交换一致。",
             1.15, 6.38, 10.85, 0.61, fill=PALE_GREEN, accent=GREEN, title_size=13, body_size=13)
    finish(slide)

    # 24 reconstruction
    slide = new_slide(prs); add_title(slide, "为什么 FNO 不能直接给出相空间涡旋：低阶矩不是 f(x,v)", section="诊断")
    add_image(slide, RESULT / "closure_fno_v3_raw_single/final_plots/phase_space_resonant_pic_vs_fno_moments.png",
              0.65, 1.35, 9.40, 5.65, caption="上：PIC 粒子分布；下：由 FNO 低阶矩构造的局域 Maxwellian")
    add_card(slide, "信息不可逆", "M₀–M₂ 只能约束 f 的前三个矩；具有不同俘获结构的分布可以共享相同低阶矩。",
             10.25, 1.55, 2.12, 1.46, fill=PALE_BLUE, accent=BLUE, title_size=13, body_size=11.5)
    add_card(slide, "图中空白", "mode-8 rollout 在 t=34.5 停止，因此 t=40/50/60 没有可重构状态。",
             10.25, 3.30, 2.12, 1.35, fill=PALE_RED, accent=RED, title_size=13, body_size=11.5)
    add_card(slide, "正确表述", "这张下排是 moment-matched Maxwellian，不是 FNO 预测的 kinetic phase space。",
             10.25, 4.95, 2.12, 1.48, fill=PALE_ORANGE, accent=ORANGE, title_size=13, body_size=11.5)
    finish(slide)

    # 25 diagnostic table
    slide = new_slide(prs); add_title(slide, "失稳归因：排除了哪些因素，还剩什么核心问题？", section="诊断")
    rows = [
        ["诊断", "结果", "结论"],
        ["dt 0.01 → 0.005", "mode-24 都停在 t=1.8", "不是时间步过大"],
        ["Ampère ↔ Poisson", "停止时间不变", "不是场求解器选择"],
        ["首个物理步给定真闭合", "无实质改善", "初值偏置不是主因"],
        ["mode 24 → 8", "t=1.8 → 34.5，但能量爆炸", "滤波只延迟失稳"],
        ["真实 PIC ∂xM₃(t), mode-24", "oracle 只到 t=8.7", "PIC 矩+截断不构成严格自治系统"],
        ["保留近全谱 mode-63", "oracle 只到 t=2.6", "高频粒子噪声会更快破坏流体积分"],
    ]
    add_table(slide, rows, 0.72, 1.42, 11.95, 4.35, widths=[3.4, 3.7, 4.85])
    add_card(slide, "核心问题 A · 数据流形外稳定性", "监督训练只见过 PIC 真轨迹；自由闭环一旦偏离，FNO 的响应没有被训练约束。",
             0.90, 6.03, 5.65, 0.86, fill=PALE_RED, accent=RED, title_size=13, body_size=12)
    add_card(slide, "核心问题 B · coarse-grained closure", "对 mode-24 的流体系统，标签应包含被截断模式反馈产生的亚网格项，而不只是过滤后的 ∂xM₃。",
             6.80, 6.03, 5.65, 0.86, fill=PALE_ORANGE, accent=ORANGE, title_size=13, body_size=12)
    finish(slide)

    # 26 matrix
    slide = new_slide(prs); add_title(slide, "复现状态矩阵：哪些结论可以说，哪些不能说")
    rows = [
        ["论文关键结论", "本项目状态", "证据"],
        ["FNO 可学习低阶矩 → 热流梯度", "✓ 复现", "Gkeyll test 0.00279；PIC raw test 0.04798"],
        ["k=.35, A=.10 进入非线性 Landau 阶段", "✓ 复现", "v=2–4.5 相空间卷曲；场能再增长"],
        ["历史不是离线高精度的必要条件", "✓ 同轨迹插值成立", "memoryless PIC corr 0.99885"],
        ["Fluid + ML 长期匹配 kinetic 场能", "✗ 未复现", "mode-24 t=1.8；mode-8 t=34.5 且发散"],
        ["PIC 可直接替代 Vlasov 训练数据", "△ 仅部分成立", "低频/线性期一致；非线性高阶矩差异显著"],
        ["当前模型可向高维直接部署", "✗ 尚不可", "一维闭环稳定性尚未解决"],
    ]
    add_table(slide, rows, 0.72, 1.42, 11.95, 4.60, widths=[4.05, 2.45, 5.45])
    add_text(slide, "最重要的研究产出不是“复现失败”，而是已经把失败压缩到可检验的两个问题：闭环离流形稳定性与谱粗粒化一致性。",
             0.98, 6.31, 11.35, 0.55, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    finish(slide)

    # 27 value/output
    slide = new_slide(prs); add_title(slide, "对我们任务的价值：已经形成可继续推进的研究资产")
    add_card(slide, "数据资产", "60 个正式 CUDA-PIC cases\nM₀–M₃ + n/u/p/q + f(x,v)\n统一存储与审计合同", 0.85, 1.55, 3.55, 1.68, fill=PALE_BLUE, accent=BLUE, title_size=16, body_size=15)
    add_card(slide, "方法资产", "Gkeyll/PIC 双数据适配\nraw/central 两套流体方程\nPoisson/Ampère、谱滤波消融", 4.62, 1.55, 3.55, 1.68, fill="F2ECF7", accent=PURPLE, title_size=16, body_size=15)
    add_card(slide, "模型资产", "v1 稳定混合闭合\nv2 history-5 研究模型\nv3 memoryless raw-moment 模型", 8.39, 1.55, 3.55, 1.68, fill=PALE_ORANGE, accent=ORANGE, title_size=16, body_size=15)
    add_card(slide, "可视化资产", "闭合时空图 / 谱噪声\n长时场能与峰值包络\n共振速度相空间与动画", 0.85, 3.72, 3.55, 1.68, fill=PALE_GREEN, accent=GREEN, title_size=16, body_size=15)
    add_card(slide, "科学认识", "PIC 与 Vlasov 的差异已量化\n无历史为何离线可行已澄清\n闭环失败机制可直接做论文", 4.62, 3.72, 3.55, 1.68, fill=PALE_RED, accent=RED, title_size=16, body_size=15)
    add_card(slide, "工程质量", "CUDA-only 正式生成\n研究/部署模型分离\n27 项回归测试通过", 8.39, 3.72, 3.55, 1.68, fill=WHITE, accent=CYAN, title_size=16, body_size=15)
    add_text(slide, "这些资产使下一轮可以直接研究稳定 closure，而不再重复数据清洗与变量语义排错。",
             1.15, 6.28, 11.05, 0.48, size=17, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    finish(slide)

    # 28 next route
    slide = new_slide(prs); add_title(slide, "下一轮建议：把“离线拟合”升级为“闭环一致的粗粒化学习”")
    next_steps = [
        ("1", "生成闭环一致标签", "对每个谱截止 K，直接计算 filtered moment 方程残差，将亚网格反馈并入 closure target。", BLUE),
        ("2", "训练离流形稳定性", "PIC 状态加物理扰动；多步自由 rollout loss；正性、能量和谱半径共同选模。", ORANGE),
        ("3", "严格因果验证", "独立 seed、独立参数、独立时间段；禁止同轨迹交错 test 作为最终结论。", PURPLE),
        ("4", "Vlasov–PIC 双基准", "用 GKEYLL 生成匹配网格轨迹，区分粒子噪声、离散耗散和模型误差。", GREEN),
        ("5", "再扩展高维", "一维纯闭环稳定并通过场能/相位验收后，再迁移到 2D/3D 与更高矩系统。", RED),
    ]
    for i, (number, title, body, color) in enumerate(next_steps):
        y = 1.40 + i * 1.03
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.92), Inches(y), Inches(0.58), Inches(0.58))
        circle.fill.solid(); circle.fill.fore_color.rgb = rgb(color); circle.line.fill.background()
        add_text(slide, number, 0.94, y + 0.08, 0.54, 0.28, size=15, color=WHITE, bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, title, 1.72, y + 0.01, 2.45, 0.30, size=16, bold=True, color=color)
        add_text(slide, body, 4.05, y + 0.01, 8.10, 0.58, size=13.5, color=INK)
    add_card(slide, "建议的下一验收门槛", "纯模型在 k=.35, A=.10 上稳定到 t=60；场能峰值包络 log10 RMSE < 0.3；总能量漂移 < 1%；再做参数外推。",
             1.12, 6.50, 11.05, 0.62, fill=PALE_ORANGE, accent=ORANGE, title_size=13, body_size=12.5)
    finish(slide)

    # Appendix divider
    slide = new_slide(prs, NAVY)
    add_text(slide, "附录", 0.85, 1.45, 3.0, 0.70, size=34, bold=True, color=ORANGE)
    add_text(slide, "完整可视化与复现实验资产", 0.88, 2.38, 7.2, 0.60, size=24, bold=True, color=WHITE)
    add_bullets(slide, ["v1 多输入/多随机种子训练曲线", "v3 三 seed 闭合时空图", "早期完整速度范围相空间与场能对齐", "纯 FNO 动画及文件清单"],
                0.95, 3.45, 6.8, 2.25, size=17, color="D3E1ED")
    add_text(slide, "主线结论不依赖附录；附录用于追溯每个图和训练选择。", 0.92, 6.35, 9.0, 0.35,
             size=14, color="9FB3C8")
    finish(slide)

    # 30 training contact
    slide = new_slide(prs); add_title(slide, "附录 A：v1 所有监督候选与随机种子训练曲线")
    add_image(slide, ASSET / "v1_training_contact.png", 0.65, 1.30, 12.00, 5.95,
              caption="nup / nupk / history / predict-q 等候选；最终选择 nupk_q_seed1")
    finish(slide)

    # 31 v3 seeds
    slide = new_slide(prs); add_title(slide, "附录 B：v3 三个随机种子均得到相近离线闭合")
    add_image(slide, ASSET / "v3_seed_contact.png", 0.65, 1.35, 12.00, 5.65,
              caption="seed 0 / 1 / 2；按 validation L2 选择 seed 1")
    add_metric(slide, "0.04449", "seed 0 best val", 2.10, 6.35, 2.25, color=BLUE, fill=PALE_BLUE)
    add_metric(slide, "0.04407", "seed 1 best val", 5.55, 6.35, 2.25, color=GREEN, fill=PALE_GREEN)
    add_metric(slide, "0.04796", "seed 2 best val", 9.00, 6.35, 2.25, color=PURPLE, fill="F2ECF7")
    finish(slide)

    # 32 earlier full phase
    slide = new_slide(prs); add_title(slide, "附录 C：早期完整速度范围相空间与场能标记")
    add_image(slide, RESULT / "closure_fno_v1/final_visualizations/group4_phase_space/pic_phase_space_evolution.png",
              0.70, 1.35, 7.15, 5.30, caption="完整 v 范围：背景占据主要动态范围")
    add_image(slide, RESULT / "closure_fno_v1/final_visualizations/group4_phase_space/phase_frames_on_energy_curve.png",
              8.05, 1.35, 4.55, 3.55, caption="相空间帧在场能演化中的对应时刻")
    add_card(slide, "改进", "后续改为背景扣除和 v=2–4.5 共振速度裁剪，才显著看见 phase-space hole。",
             8.15, 5.28, 4.20, 1.05, fill=PALE_ORANGE, accent=ORANGE, title_size=13, body_size=12)
    finish(slide)

    # 33 pure animation and strong case
    slide = new_slide(prs); add_title(slide, "附录 D：纯 FNO 相空间动画与强案例场能")
    add_image(slide, RESULT / "closure_fno_v2/pure_fno_longtime/k0p400_a0p050_pic_vs_pure_fno_maxwellian.gif",
              0.75, 1.35, 6.10, 4.95, caption="GIF：放映模式下播放 PIC 与纯 FNO moment-matched Maxwellian")
    add_image(slide, RESULT / "closure_fno_v2/pure_fno_longtime/k0p350_a0p100_pure/k0p350_a0p100_field_energy.png",
              7.05, 1.35, 5.55, 4.15, caption="k=.35, A=.10 强案例场能与停止点")
    add_card(slide, "说明", "相空间动画下排仍不是 kinetic f 预测，只是方便观察低阶矩隐含的局部 Maxwellian 演化。",
             7.20, 5.78, 5.15, 0.85, fill=PALE_BLUE, accent=BLUE, title_size=13, body_size=12)
    finish(slide)

    # 34 reproducibility
    slide = new_slide(prs); add_title(slide, "附录 E：可复现实验文件与质量状态")
    rows = [
        ["资产", "路径/状态"],
        ["v1 工程模型", "models/closure/closure_fno_pic_v1.pt（当前部署默认）"],
        ["v2 history-5", "models/closure/closure_fno_pic_v2_history5.pt（研究模型）"],
        ["v3 raw single", "models/closure/closure_fno_pic_v3_raw_single_research.pt（拒绝部署）"],
        ["v3 指标", "results/closure_fno_v3_raw_single/final_report/study_summary.json"],
        ["中文实验报告", "docs/reports/PIC_RAW_MOMENT_FNO_EXPERIMENT_zh-CN.md"],
        ["回归测试", "27/27 passed；新增 raw moment derivative/filter/split 测试"],
    ]
    add_table(slide, rows, 0.72, 1.42, 11.95, 4.55, widths=[3.1, 8.85])
    add_card(slide, "模型治理", "v3 checkpoint 已冻结并校验 SHA-256，但 deployment_approved=false，不覆盖 v1 默认模型。",
             1.05, 6.20, 11.20, 0.70, fill=PALE_GREEN, accent=GREEN, title_size=13, body_size=12.5)
    finish(slide)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    manifest = {
        "presentation": str(OUT),
        "slide_count": len(prs.slides),
        "size_bytes": OUT.stat().st_size,
        "generated_assets": sorted(str(path) for path in ASSET.iterdir()),
    }
    (OUT.with_suffix(".manifest.json")).write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return OUT


if __name__ == "__main__":
    print(build_presentation())
