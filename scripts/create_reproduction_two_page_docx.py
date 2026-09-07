#!/usr/bin/env python3
"""Create a concise two-page Chinese reproduction note."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/reports/GKEYLL_FNO_REPRODUCTION_TWO_PAGE_zh-CN.docx"
R8 = ROOT / "results/gkeyll_round8/final_report"
R9 = ROOT / "results/gkeyll_round9/final_report"

FIELD = R8 / "paper_field_energy_truth_hp_fno.png"
MOMENTS = R8 / "paper_moments_truth_hp_fno.png"
TEMPORAL = R9 / "single_case_temporal_extrapolation.png"
CROSS = R9 / "cross_case_generalization.png"

NAVY = "17365D"
BLUE = "2F75B5"
TEAL = "0F6B78"
LIGHT_BLUE = "DDEBF7"
LIGHT_TEAL = "DDEFEF"
LIGHT_GRAY = "F3F5F7"
AMBER = "FFF2CC"
GREEN = "E2F0D9"
GRAY = RGBColor(90, 98, 108)


def font(run, size=9, bold=False, color=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, value=75):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("第 ")
    font(r, 7.5, color="777777")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.extend((begin, instr, end))
    r = paragraph.add_run(" / 2")
    font(r, 7.5, color="777777")


def setup(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(0.85)
    section.bottom_margin = Cm(0.85)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    section.header_distance = Cm(0.35)
    section.footer_distance = Cm(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(8.8)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.05

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("Huang et al. (2025) FNO 闭合复现：问题发现与解决")
    font(r, 7.5, color="777777")
    add_page_field(section.footer.paragraphs[0])


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    font(r, 18, True, NAVY)
    if subtitle:
        r = p.add_run("  |  " + subtitle)
        font(r, 9.2, False, TEAL)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), BLUE)
    borders.append(bottom)
    p_pr.append(borders)


def add_callout(doc, lead, body, fill=GREEN):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, fill)
    margins(c, 95)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(lead + "  ")
    font(r, 9.2, True, TEAL if fill == GREEN else NAVY)
    r = p.add_run(body)
    font(r, 8.8)


def add_table(doc, headers, rows, widths=None, font_size=8.0):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        t.rows[0].cells[i].text = value
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
    for ridx, row in enumerate(t.rows):
        no_split(row)
        for cidx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell, 65)
            if widths:
                cell.width = Cm(widths[cidx])
            if ridx == 0:
                shade(cell, NAVY)
            elif ridx % 2 == 0:
                shade(cell, LIGHT_GRAY)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    font(r, font_size, ridx == 0, "FFFFFF" if ridx == 0 else None)
    return t


def add_section_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    font(r, 10.5, True, TEAL)


def add_figure_cell(cell, path, width, caption):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Inches(width))
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(caption)
    font(r, 7.5, False, "555555")


def page_one(doc):
    add_title(doc, "FNO 热流闭合复现：问题发现与解决", "两页内部说明")
    add_callout(
        doc,
        "一句话结论：",
        "按论文颗粒度对齐 Gkeyll 数据、中心矩、∂xq 标签、三矩方程和 RK4 闭环后，我们在同案例/同轨迹上复现了非线性场能回升；"
        "目前尚未解决的是未见时间与未见参数下的长期闭环泛化。",
    )

    add_section_label(doc, "1  主要问题是怎样发现并解决的")
    rows = [
        ("数据源与非线性不对齐", "PIC 高阶矩噪声大；弱扰动多数只出现相混合条纹", "先用论文 Gkeyll 锚点 k=.35,A=.10、t=0–40 的完整强非线性轨迹复现；PIC 留作域迁移", "同轨迹出现阻尼—最低点—回升"),
        ("物理量语义不对齐", "原始 M2/M3 与中心 p/q 混用；预测 q 再求导会放大噪声", "严格使用中心矩 p=M2−nu²、q=M3−3uM2+2nu³，并直接监督压力方程所需的 ∂xq", "闭合 rel-L2=0.00458，相关系数=0.99999"),
        ("PDE 与调用粒度不对齐", "Poisson/Ampère、primitive/守恒形式及每整步一次闭合混杂", "对齐 primitive 三矩+Ampère；每个 RK4 stage 调 FNO；第一步使用真实初始闭合", "真值 closure oracle 场能误差约 10⁻³，排除 PDE 写错"),
        ("网络结构与频谱不对齐", "GroupNorm/GELU 和高模输出导致闭环敏感", "改为论文明确的 4 层+ReLU、无 GroupNorm；扫描 8/12/16/24 后部署 mode 8", "Round6 相对旧结构改善 58.1%；最佳场能误差 0.0379"),
        ("图好看但评估口径过强", "完整轨迹训练后仍在同一案例验证；最佳 seed 掩盖波动", "同时报告 3-seed 中位数，并补做 t<24 时间外推和整案例留出", "同轨迹中位 0.1788；严格泛化升至 0.44–0.96"),
    ]
    add_table(doc, ["发现的问题", "诊断证据", "解决办法", "解决后的结果"], rows, [4.0, 6.5, 9.2, 6.4], 7.7)

    add_section_label(doc, "2  与论文的物理合同是否对齐")
    rows = [
        ("已对齐", "1D1V Gkeyll Vlasov–Ampère；固定中和离子；k=.35,A=.10；Nx=Nv=64；v∈[−6,6]；t≤40；中心矩 {n,u,p}→∂xq；训练 Δt=.005；RK4 闭环 Δt=.002。"),
        ("工程选择", "论文正文未披露 width/modes；我们训练用 128/32，部署额外采用 mode 8。严格时间外推与整案例留出是我们新增的评估，不是论文主图协议。"),
    ]
    t = add_table(doc, ["结论", "内容"], rows, [3.0, 23.1], 7.8)
    shade(t.rows[1].cells[0], GREEN)
    shade(t.rows[2].cells[0], AMBER)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("核心判断：")
    font(r, 8.5, True, NAVY)
    r = p.add_run("改善主要来自数据合同、物理合同、数值合同和评估合同同时对齐，而不是单纯把 FNO 做大。")
    font(r, 8.5)


def page_two(doc):
    doc.add_page_break()
    add_title(doc, "关键结果与当前边界", "四张图回答“解决了什么、还差什么”")

    grid = doc.add_table(rows=2, cols=2)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for row in grid.rows:
        no_split(row)
        for cell in row.cells:
            margins(cell, 25)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    add_figure_cell(
        grid.cell(0, 0), FIELD, 4.95,
        "① 同轨迹闭环：FNO 与真值近乎重合；HP 不能保持非线性回升（FNO 0.0379，HP 1.1773）",
    )
    add_figure_cell(
        grid.cell(0, 1), TEMPORAL, 4.95,
        "② 时间外推：只训练 t<24 后，未见非线性段闭合误差突增；场能误差中位 0.4443",
    )
    add_figure_cell(
        grid.cell(1, 0), MOMENTS, 4.72,
        "③ 同轨迹低阶矩：FNO 保持 n、u、p 的幅值和相位，HP 中后期持续衰减",
    )
    add_figure_cell(
        grid.cell(1, 1), CROSS, 4.72,
        "④ 整案例留出：瞬时闭合相关系数约 0.997，但 20,000 步反馈后仍出现幅值和相位漂移",
    )

    add_callout(
        doc,
        "最终结论：",
        "已经解决的是“严格按论文单案例训练时，FNO 闭合能否正确嵌入流体方程并优于 HP”；"
        "尚未解决的是“该闭合能否外推到后期未见状态和未见参数案例”。下一步应优先扩充强非线性案例、按整轨迹划分数据，"
        "并用短闭环验证/偏移状态训练约束能量交换和相位，而不是继续只优化逐帧 ∂xq 误差。",
        AMBER,
    )


def main():
    for path in (FIELD, MOMENTS, TEMPORAL, CROSS):
        if not path.exists():
            raise FileNotFoundError(path)
    doc = Document()
    setup(doc)
    doc.core_properties.title = "FNO 热流闭合复现：问题发现与解决（两页版）"
    doc.core_properties.subject = "Huang et al. (2025) 复现内部说明"
    doc.core_properties.author = "landau-damping-surrogate-standardized 项目组"
    page_one(doc)
    page_two(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
